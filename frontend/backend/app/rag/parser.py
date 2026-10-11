"""Text and tables out of a document: PDF, XLSX, CSV, DOCX, TXT or Markdown (RAG.md §4)."""
import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import PurePath
from typing import Iterable, Literal, Optional
from xml.etree import ElementTree

from openpyxl import load_workbook
from pypdf import PdfReader

DocType = Literal['market_report', 'freight', 'agronomy', 'methodology', 'contract', 'other']
SUPPORTED = ('.pdf', '.xlsx', '.csv', '.docx', '.txt', '.md')
MAX_XML_BYTES = 50_000_000  # a .docx is a zip: its text is not allowed to unpack to more than this


class UnsupportedFile(ValueError):
    """Not a format that can be indexed."""


class UnreadableFile(ValueError):
    """The right extension, but damaged or without any text."""


@dataclass
class Block:
    """A stretch of text, or a table (header row first), with where it sits in the document."""
    text: str = ''
    rows: list[list[str]] = field(default_factory=list)
    page: Optional[int] = None  # PDF page
    section: Optional[str] = None  # heading it falls under, or the sheet of a workbook


@dataclass
class Parsed:
    blocks: list[Block]
    pages: Optional[int] = None  # PDF only
    title: Optional[str] = None  # when the document names itself (first heading of a Markdown file)


def _cell(v: object) -> str:
    if v is None:
        return ''
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == time.min else v.isoformat(sep=' ')
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float):
        return f'{v:.10g}'  # 25.0 → 25, and no binary noise such as 0.30000000000000004
    return str(v).strip()


def _table(rows: Iterable[Iterable[object]]) -> list[list[str]]:
    """Cells as text, without the empty rows and the empty columns on the right."""
    cells = [[_cell(c) for c in r] for r in rows]
    cells = [r for r in cells if any(r)]
    width = max((i + 1 for r in cells for i, c in enumerate(r) if c), default=0)
    return [(r + [''] * width)[:width] for r in cells]


def _decode(data: bytes) -> str:
    try:
        return data.decode('utf-8-sig')
    except UnicodeDecodeError:
        return data.decode('latin-1')


def _pdf(data: bytes) -> Parsed:
    reader = PdfReader(io.BytesIO(data))
    blocks = []
    for number, page in enumerate(reader.pages, 1):
        # A PDF breaks the line wherever the page is full: only blank lines are paragraph ends.
        text = re.sub(r'(?<!\n)[ \t]*\n(?!\s*\n)', ' ', page.extract_text() or '').strip()
        if text:
            blocks.append(Block(text=text, page=number))
    return Parsed(blocks, pages=len(reader.pages))


def _xlsx(data: bytes) -> Parsed:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    return Parsed([Block(rows=rows, section=sheet.title) for sheet in workbook.worksheets if (rows := _table(sheet.iter_rows(values_only=True)))])


def _csv(data: bytes) -> Parsed:
    text = _decode(data)
    header = text.lstrip().split('\n', 1)[0]
    delimiter = max(',;\t', key=header.count)  # the separator the header line uses most; a comma when it has none
    rows = _table(csv.reader(io.StringIO(text), delimiter=delimiter))
    return Parsed([Block(rows=rows)] if rows else [])


W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def _docx(data: bytes) -> Parsed:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if archive.getinfo('word/document.xml').file_size > MAX_XML_BYTES:
            raise ValueError('document.xml is too large')
        body = ElementTree.fromstring(archive.read('word/document.xml')).find(f'{W}body')
    text_of = lambda el: ''.join(t.text or '' for t in el.iter(f'{W}t')).strip()
    blocks: list[Block] = []
    section = None
    for el in body if body is not None else ():
        if el.tag == f'{W}tbl':
            rows = _table([text_of(cell) for cell in row.iter(f'{W}tc')] for row in el.iter(f'{W}tr'))
            if rows:
                blocks.append(Block(rows=rows, section=section))
        elif el.tag == f'{W}p' and (text := text_of(el)):
            style = el.find(f'{W}pPr/{W}pStyle')
            if style is not None and re.match('heading|title', style.get(f'{W}val', ''), re.I):
                section = text
            else:
                blocks.append(Block(text=text, section=section))
    return Parsed(blocks)


def _txt(data: bytes) -> Parsed:
    text = _decode(data).strip()
    return Parsed([Block(text=text)] if text else [])


def _markdown(data: bytes) -> Parsed:
    """One block per heading. The section is the path of headings above the text, without the document title."""
    blocks: list[Block] = []
    path: list[str] = []
    lines: list[str] = []
    title = None
    fenced = False

    def flush() -> None:
        text = '\n'.join(lines).strip()
        if text:
            # The title is on every citation already: it is not repeated as the section.
            blocks.append(Block(text=text, section=' › '.join(path[1:] if path[:1] == [title] else path) or None))
        lines.clear()

    for line in _decode(data).splitlines():
        if line.lstrip().startswith('```'):
            fenced = not fenced
        heading = None if fenced else re.match(r'(#{1,6})\s+(.+?)\s*#*\s*$', line)
        if heading:
            flush()
            name = re.sub(r'[`*_]', '', heading[2]).strip()
            path[len(heading[1]) - 1:] = [name]
            if title is None and len(heading[1]) == 1:
                title = name
        else:
            lines.append(line)
    flush()
    return Parsed(blocks, title=title)


READERS = {'.pdf': _pdf, '.xlsx': _xlsx, '.csv': _csv, '.docx': _docx, '.txt': _txt, '.md': _markdown}


def parse(name: str, data: bytes) -> Parsed:
    ext = PurePath(name).suffix.lower()
    if ext == '.xls':
        raise UnsupportedFile('Legacy .xls files cannot be indexed: save the file as .xlsx')
    if ext not in READERS:
        raise UnsupportedFile(f'Only {", ".join(SUPPORTED)} files can be indexed')
    try:
        parsed = READERS[ext](data)
    except Exception as e:  # each reader has its own errors for a damaged file (bad zip, bad xref, bad XML…)
        raise UnreadableFile(f'{name} could not be read as a {ext} file') from e
    if not parsed.blocks:
        raise UnreadableFile(f'No text found in {name}' + (' (a scanned PDF needs OCR first)' if ext == '.pdf' else ''))
    return parsed


# First match wins on a tie, so the more specific kinds come first.
KEYWORDS: dict[str, str] = {
    'freight': r'freight|vessel|shipping|handysize|supramax|charter|cargo',
    'agronomy': r'agro|trial|yield|soil|application rate|field stud|hectare|\bt/ha\b',
    'contract': r'contract|agreement|offtake|incoterm|clause',
    'methodology': r'methodolog|polic(y|ies)|guideline|specification|formula|value shar|carbon|cbam|scope [123]',
    'market_report': r'argus|\bcru\b|price|\bfob\b|\bcfr\b|market|benchmark|quote',
}


def classify(name: str, text: str) -> DocType:
    """Kind of document, from the words in its filename (which count most) and in its first lines."""
    name, text = re.sub(r'[-_.]+', ' ', name.lower()), text[:4000].lower()  # value-sharing-note.md says "value sharing"
    score = {kind: 5 * len(re.findall(words, name)) + len(re.findall(words, text)) for kind, words in KEYWORDS.items()}
    best = max(score, key=score.get)
    return best if score[best] else 'other'  # type: ignore[return-value]


def date_in_name(name: str) -> Optional[date]:
    """dd-mm-yyyy or yyyy-mm-dd anywhere in the filename."""
    found = [(m[3], m[2], m[1]) for m in re.finditer(r'(\d{2})-(\d{2})-(\d{4})', name)]
    found += [(m[1], m[2], m[3]) for m in re.finditer(r'(\d{4})-(\d{2})-(\d{2})', name)]
    for y, m, d in found:
        try:
            return date(int(y), int(m), int(d))
        except ValueError:
            continue
    return None
